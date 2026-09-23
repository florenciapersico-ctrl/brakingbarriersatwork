# Genera el documento de revisión "Master Delivery Map · 20 sesiones 1:1"
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BROWN, DARK, TERRA, MUTED = "81552E", "3F1E0D", "C68C6C", "8A6A55"
TINT, CARD = "FBF2D0", "F8F5EE"
TITLE, BODY = "Source Serif Pro", "Open Sans"

S = [
 dict(n=1, sem=1, etapa="Onboarding", nombre="Onboarding de la alumna",
  core="Instalar su sistema de trabajo y fijar su punto de partida.",
  plat="ONB 01–13 + recurso «Tu recorrido AT WORK».",
  res="Sabe cómo funciona AT WORK, tiene su Starting Point revisado, un foco inicial definido en una frase y su primera situación real elegida.",
  ev="Starting Point completo (autoevaluación + audio o video espontáneo de 2–3 min) · checklist de ONB 13 completo · foco inicial escrito · horario 1:1 agendado.",
  foco="instalar tu sistema de trabajo y definir desde dónde partes", puede="empezar el Módulo 1 sabiendo cuál es tu prioridad",
  micro=["El recorrido en 6 etapas, con el recurso «Tu recorrido AT WORK».",
         "El ciclo de cada semana: situación real → producir → feedback → retry → evidencia.",
         "Qué va en cada espacio (Skool, AT WORK Community, Drive, sesión 1:1) y cómo funciona el Weekly Report."],
  app="Escuchan juntas un fragmento de su audio del Starting Point. Detectas 1–2 patrones de alto impacto y, a partir de eso, definen el foco inicial. Después eligen la situación real que va a traer a la sesión 2. Ej.: Belén, el onboarding de un proveedor nuevo; Sol, el status de su proyecto; Virginia, el update al CEO.",
  retry="Se presenta (rol, equipo, qué hace) en 60 segundos, sin guion. Le das una sola indicación y lo repite.",
  next="Completar lo que falte del checklist de ONB 13, subir 3–5 materiales reales, reservar sus bloques de práctica, ver M1.1–M1.3 y enviar su primer Weekly Report.",
  ya="En esta sesión todavía no se enseña inglés: es instalación. Si todo el checklist está completo, dedica más tiempo al audio del Starting Point."),

 dict(n=2, sem=2, etapa="Activate Your English", nombre="Situación real → producir primero",
  core="Convertir una situación real de su trabajo en práctica y producir antes de preparar (Role Play Method).",
  plat="M1.1 My Professional Context · M1.2 Practice System · M1.3 Role Play Method · M1.7 Bots.",
  res="Arma un role play a partir de una situación real y habla antes de prepararse.",
  ev="Role play grabado (intento 1 + retry) sobre su situación real · lista de 3–5 situaciones recurrentes de su trabajo en 01 · MI CONTEXTO PROFESIONAL.",
  foco="convertir tu situación real en práctica", puede="practicar tu propio trabajo entre sesiones, con o sin bot",
  micro=["Por qué producir primero: el inglés que ya tienes se activa usándolo, no repasándolo.",
         "Cómo se arma un role play: quién es el interlocutor, qué necesita, qué quieres lograr, 2–3 minutos.",
         "Cómo usar los bots para repetir esa situación entre sesiones."],
  app="Role play en vivo, tú haces del interlocutor. Belén le explica al proveedor los pasos del onboarding; Sol le da el status a su manager; Virginia abre el update al CEO.",
  retry="Mismo role play con un solo ajuste del feedback. Presión: el interlocutor hace una pregunta inesperada.",
  next="Dos role plays con bot sobre situaciones de su lista, audios guardados en 04 · MIS PRÁCTICAS. Weekly Report.",
  ya="Sube la presión: interlocutor con más poder, menos tiempo o una situación que evita."),

 dict(n=3, sem=3, etapa="Activate Your English", nombre="Narrative Flow",
  core="Contar una situación de trabajo con un hilo claro, sin frases sueltas.",
  plat="M1.4 Narrative Flow.",
  res="Cuenta qué pasó o qué está pasando en su trabajo con inicio, desarrollo y cierre, conectando las ideas.",
  ev="Audio de 1–2 min narrando una situación real con secuencia clara y conectores · comparación intento 1 vs retry.",
  foco="contar lo que pasa en tu trabajo con un hilo claro", puede="contar una situación de principio a fin sin perderte",
  micro=["Estructura: contexto → qué pasó → qué hiciste o qué se decidió → resultado o próximo paso.",
         "Conectores: first, then, after that, because, so, in the end.",
         "Pasado para lo que ocurrió, presente para cómo están las cosas hoy."],
  app="Cuenta lo que pasó esta semana. Belén: cómo fue el primer contacto con el proveedor. Sol: qué cambió en el proyecto desde el último update. Virginia: la semana de su equipo en 90 segundos.",
  retry="La misma historia en menos tiempo (de 2 minutos a 90 segundos). Presión: la interrumpes con «Wait, why did that happen?» y retoma el hilo.",
  next="Narrar una situación de trabajo por día en un audio de 1 minuto (o con bot). Weekly Report.",
  ya="Pasa a narrar un problema: qué lo causó y qué consecuencias tuvo."),

 dict(n=4, sem=4, etapa="Activate Your English", nombre="Color Method + ciclo completo",
  core="Cerrar el ciclo completo: producir → analizar → mejorar → practicar → retry → aplicar.",
  plat="M1.5 Color Method · M1.6 Complete Practice Cycle / Retry.",
  res="Revisa su propia producción, identifica el lenguaje que le faltó y lo recupera en el retry sin mirar notas.",
  ev="Producción original + retry con al menos una mejora visible · lenguaje nuevo registrado con Color Method · primer ciclo completo guardado en 04 · MIS PRÁCTICAS.",
  hito="Primer ciclo completo. Pregunta clave: ¿puede practicar sola entre sesiones? Si no, refuérzalo antes de avanzar a la etapa 2.",
  foco="cerrar tu primer ciclo completo de práctica", puede="practicar sola, con método, entre sesiones",
  micro=["Color Method: marcar el lenguaje relevante para pasarlo de reconocerlo a recuperarlo y usarlo (con los colores de la clase M1.5).",
         "El ciclo completo del Módulo 1, de punta a punta.",
         "Qué es un buen retry: la misma tarea, una mejora concreta y sin leer."],
  app="Retoma la situación de las sesiones 2–3 para poder comparar. Graba, marca con Color Method el lenguaje que necesitaba y hace el retry.",
  retry="Retry sin notas. Presión: una pregunta de seguimiento sobre lo que acaba de decir.",
  next="Un ciclo completo por su cuenta con otra situación de su lista.",
  ya="Que haga el ciclo con una situación nueva en la sesión y tú solo observes."),

 dict(n=5, sem=5, etapa="Communicate in Real Time", nombre="Propósito y mensaje clave",
  core="Saber para qué habla y cuál es su mensaje clave; construir frases claras y cortas.",
  plat="M2.1 Build Clear Sentences.",
  res="Antes de hablar, define su propósito y su mensaje clave en una frase, y lo dice con oraciones cortas y completas.",
  ev="Propósito y mensaje clave escritos en una línea cada uno · audio donde el mensaje clave aparece en los primeros 20 segundos · menos frases largas encadenadas.",
  foco="saber qué quieres lograr y decirlo con claridad", puede="decir tu mensaje clave en una frase",
  micro=["Tres preguntas antes de hablar: ¿quién escucha?, ¿qué necesita saber o decidir?, ¿cuál es mi mensaje en una frase?",
         "Oración clara: sujeto + verbo + información. Una idea por frase.",
         "Chunks: «The key point is…», «What I need from you is…», «The main thing is…»."],
  app="Belén: el mensaje clave para el proveedor («To start working with us, you need to complete three steps this week»). Sol: el mensaje clave de su status. Virginia: lo único que el CEO tiene que saber.",
  retry="La misma intervención, empezando por el mensaje clave y con frases cortas. Presión: «So what do you need from me?».",
  next="Antes de cada reunión de la semana, escribir propósito + mensaje clave. Grabar un ejemplo.",
  ya="Que lo haga sin escribir: 10 segundos para pensar y habla."),

 dict(n=6, sem=6, etapa="Communicate in Real Time", nombre="Main Point First",
  core="Organizar y priorizar: el punto principal primero, con Message Map.",
  plat="M2.3 Organize & Prioritize.",
  res="Da un update claro en 1–2 minutos: primero el punto principal, después el soporte y cierra con un próximo paso.",
  ev="Message Map de una situación real · update grabado de 1–2 min que empieza por el punto principal y cierra con una acción.",
  foco="empezar por lo importante", puede="dar un update claro en 1–2 minutos",
  micro=["Message Map: mensaje principal → 2–3 puntos de soporte → próximo paso.",
         "BLUF (Bottom Line Up Front).",
         "Por qué en inglés profesional el contexto va después del punto, no antes."],
  app="Sol: el status de su proyecto. Virginia: el update al CEO. Belén: un resumen para su equipo sobre cómo va el onboarding del proveedor.",
  retry="El mismo update con Message Map, sin leer. Presión: «I only have one minute».",
  next="Usar el Message Map antes de una reunión real y contar en el Weekly Report qué pasó.",
  ya="Que arme el Message Map mentalmente en 30 segundos, sin papel."),

 dict(n=7, sem=7, etapa="Communicate in Real Time", nombre="Frameworks flexibles",
  core="Elegir la estructura adecuada según la situación.",
  plat="M2.3 Organize & Prioritize (frameworks).",
  res="Elige y usa la estructura adecuada para un update, un problema o una recomendación.",
  ev="La misma situación dicha con dos frameworks distintos · puede explicar cuál eligió y por qué.",
  foco="elegir la estructura según la situación", puede="organizar un update, un problema o una recomendación sin guion",
  micro=["STATUS → DELTA → WHY → ACTION → OUTLOOK, para updates.",
         "Situation → Issue → Action, para problemas.",
         "Point → Evidence → Implication, para recomendaciones.",
         "Una estructura es un mapa, no un guion."],
  app="Sol: un retraso en el proyecto (Situation → Issue → Action). Virginia: una recomendación al CEO (Point → Evidence → Implication). Belén: el status del proveedor (STATUS → DELTA…).",
  retry="Repite con el framework que mejor le funcionó. Presión: cambia el interlocutor (de su manager a alguien de otra área).",
  next="Identificar qué situación de su semana pide cada framework y practicar una con bot.",
  ya="Presión: se le cambia el objetivo a mitad de camino y tiene que reorganizar."),

 dict(n=8, sem=8, etapa="Communicate in Real Time", nombre="Keep Speaking",
  core="Seguir hablando cuando no aparece la palabra: parafrasear, describir, simplificar.",
  plat="M2.2 Keep Speaking When the Word Doesn't Come.",
  res="Cuando no encuentra una palabra, la rodea y sigue, sin frenar ni pasar al español.",
  ev="Producción donde resuelve 2–3 bloqueos parafraseando · menos silencios largos y reinicios.",
  foco="seguir hablando aunque falte una palabra", puede="resolver un bloqueo sin frenar",
  micro=["Estrategias: describir la función («the thing we use to…»), dar un ejemplo, usar una palabra más general, decir lo contrario («it's not…»), simplificar la idea.",
         "Chunks: «What I mean is…», «It's a kind of…», «Let me put it another way»."],
  app="Explica un término de su trabajo sin poder usar la palabra clave. Belén: un requisito de compliance del proveedor. Sol: una dependencia técnica. Virginia: un indicador financiero.",
  retry="La misma explicación con otras palabras prohibidas. Presión: «Sorry, I don't understand. Can you explain it differently?».",
  next="Juego de palabras prohibidas con bot, con tres términos de su trabajo.",
  ya="Que explique lo mismo a alguien que no es de su área."),

 dict(n=9, sem=9, etapa="Communicate in Real Time", nombre="Repair & Continue + manejar la interacción",
  core="Ganar tiempo, pedir aclaración, recuperar el hilo y responder preguntas e interrupciones.",
  plat="M2.4 Repair & Continue · M2.5 Manage the Interaction.",
  res="Ante una pregunta, una interrupción o un error propio, se recupera y sigue sin reiniciar la explicación.",
  ev="Role play con interrupciones donde gana tiempo, pide aclaración y retoma el hilo · responde las preguntas con una respuesta directa primero.",
  foco="recuperarte y seguir cuando te interrumpen o te preguntan", puede="manejar preguntas sin perder el hilo",
  micro=["Ganar tiempo: «That's a good question. Let me think…».",
         "Pedir aclaración: «Do you mean…?».",
         "Retomar: «Going back to…», «As I was saying…». Autocorregirse sin frenar: «Sorry, I mean…».",
         "Responder preguntas: respuesta directa + una razón."],
  app="Ronda de preguntas sobre su situación. Virginia: preguntas incómodas del CEO. Sol: un stakeholder impaciente. Belén: un proveedor que no entiende un paso.",
  retry="La misma ronda de preguntas. Presión: dos preguntas seguidas o una objeción.",
  next="Anticipar tres preguntas difíciles de una reunión real y practicarlas con bot.",
  ya="Objeción con desacuerdo: tiene que sostener su postura."),

 dict(n=10, sem=10, etapa="Communicate in Real Time", nombre="Clara, concisa y específica + revisión de mitad",
  core="Integrar la etapa en una situación real completa: clara, concisa y específica.",
  plat="M2.6 Workplace Practice Library + la Workplace Guide que corresponda.",
  res="Maneja una situación real completa (update + preguntas) de forma clara, concisa y específica.",
  ev="Situación completa grabada y comparada con el Starting Point y con la sesión 2 · lista de 2–3 patrones lingüísticos prioritarios para Polish.",
  hito="Revisión de mitad de programa. Comparan con el Starting Point: qué ya hace sola y qué patrones lingüísticos todavía limitan su claridad. Eso define el foco de las sesiones 11–16.",
  foco="integrar todo lo de esta etapa y mirar cuánto avanzaste", puede="manejar una situación real completa y saber qué pulir",
  micro=["Clara, concisa y específica: cortar contexto innecesario, usar números y fechas concretas, cerrar con acción, responsable y plazo.",
         "La Workplace Guide que corresponda a su situación."],
  app="La situación más importante de su mes. Belén: el kickoff completo con el proveedor. Sol: el status con preguntas. Virginia: el update al CEO con Q&A.",
  retry="Una versión más corta y más específica. Presión: preguntas sobre números o fechas.",
  next="Confirmar juntas los 2–3 patrones prioritarios para las semanas 11–16.",
  ya="—"),

 dict(n=11, sem=11, etapa="Polish Your English", nombre="Verb System I",
  core="Tiempos verbales para lo que pasó, lo que está pasando y lo habitual.",
  plat="M3.1 Verb System.",
  res="Usa el tiempo correcto para hablar de lo que ya pasó, lo que está pasando y cómo funcionan las cosas en su trabajo.",
  ev="Update donde su patrón prioritario de tiempos verbales sale bien en la mayoría de los casos · comparación intento 1 vs retry.",
  foco="usar bien los tiempos verbales en tus updates", puede="contar qué hicieron, qué están haciendo y cómo funciona",
  micro=["La línea de tiempo de un update: what we did (past simple) / what we're doing (present continuous) / how it usually works (present simple).",
         "Su error más frecuente, según la revisión de la sesión 10."],
  app="Sol: qué hicieron, qué están haciendo y cómo sigue. Belén: cómo funciona normalmente el onboarding y qué está pasando con este proveedor. Virginia: resultados de la semana.",
  retry="El mismo update, atenta a los verbos. Presión: «When did that happen?».",
  next="Grabar el update semanal y marcar los verbos con Color Method.",
  ya="Pasa directo a Verb System II."),

 dict(n=12, sem=12, etapa="Polish Your English", nombre="Verb System II",
  core="Present perfect y futuro: resultados hasta ahora, planes y compromisos.",
  plat="M3.1 Verb System.",
  res="Distingue resultados hasta ahora de hechos terminados, y habla de planes y compromisos con la forma adecuada.",
  ev="Update con «so far / already / yet» y próximos pasos en futuro, bien usados.",
  foco="hablar de resultados y de próximos pasos", puede="decir qué está hecho, qué falta y qué viene",
  micro=["«We've completed…» (resultado hasta ahora) vs «We completed it on Monday» (hecho con fecha).",
         "Futuro: plan («We're going to…»), algo agendado («We're meeting them on Friday»), compromiso en el momento («I'll send it today»)."],
  app="Virginia: resultados del trimestre hasta ahora y próximos pasos. Belén: qué ya completó el proveedor y qué falta. Sol: avances y próximos hitos del proyecto.",
  retry="El mismo update. Presión: «What's still pending?».",
  next="Update semanal con «so far» y tres próximos pasos en futuro.",
  ya="Que sume condiciones: «If they approve it, we'll…»."),

 dict(n=13, sem=13, etapa="Polish Your English", nombre="Questions & Word Order",
  core="Hacer preguntas claras y ordenar bien la frase.",
  plat="M3.2 Questions & Word Order.",
  res="Formula preguntas directas e indirectas correctas y mantiene el orden de la frase en oraciones largas.",
  ev="Cinco preguntas reales de su trabajo, correctas · role play donde ella dirige con preguntas.",
  foco="preguntar bien y ordenar tus frases", puede="conseguir la información que necesitas con preguntas claras",
  micro=["Estructura: (question word) + auxiliar + sujeto + verbo.",
         "Preguntas indirectas, más amables y sin inversión: «Could you tell me when…?», «Do you know if…?».",
         "Orden de la frase: sujeto + verbo + objeto + lugar + tiempo; el adjetivo va antes del sustantivo."],
  app="Belén le pregunta al proveedor por documentación pendiente. Sol le pide información a otro equipo. Virginia hace preguntas de seguimiento a su equipo.",
  retry="Las mismas preguntas. Presión: el interlocutor responde vago y ella tiene que repreguntar.",
  next="Preparar las preguntas de una reunión real y usarlas.",
  ya="Que conduzca una reunión corta solo con preguntas."),

 dict(n=14, sem=14, etapa="Polish Your English", nombre="Modals & Functional Accuracy",
  core="Pedir, sugerir, recomendar, marcar obligación y suavizar con modales.",
  plat="M3.3 Modals.",
  res="Usa modales para pedir, recomendar, marcar obligación y suavizar sin sonar brusca ni ambigua.",
  ev="Un pedido, una recomendación y un límite en la misma interacción, con el modal adecuado.",
  foco="pedir y recomendar con el tono justo", puede="ser firme y cordial a la vez",
  micro=["Pedir: could / would you. Recomendar: should / I'd recommend.",
         "Obligación: need to / have to / must. Posibilidad: might / could.",
         "Suavizar: «It might be better to…», «Would it be possible to…?»."],
  app="Belén: pedirle al proveedor que cumpla un plazo, firme pero cordial. Sol: pedir más recursos. Virginia: recomendarle una decisión al CEO.",
  retry="La misma conversación. Presión: la otra persona dice que no.",
  next="Usar un pedido y una recomendación en una situación real y contarlo en el Weekly Report.",
  ya="Presión: negociar una alternativa después del no."),

 dict(n=15, sem=15, etapa="Polish Your English", nombre="Causas, condiciones y consecuencias + patrones de alta frecuencia",
  core="Explicar por qué pasa algo, qué pasa si, y qué implica; corregir sus patrones de alta frecuencia.",
  plat="M3.4 Causes, Conditions & Consequences · M3.5 High-Frequency Precision Patterns.",
  res="Explica causas, condiciones y consecuencias con conectores precisos, y corrige su patrón de alta frecuencia prioritario.",
  ev="Explicación de un problema real con because / due to, if / unless, so / as a result · su error de alta frecuencia baja visiblemente en el retry.",
  foco="explicar causas y consecuencias con precisión", puede="explicar un problema y su impacto con claridad",
  micro=["Causa: because / because of / due to.",
         "Condición: if / unless / as long as (if + presente, will).",
         "Consecuencia: so / as a result / which means.",
         "Su patrón de alta frecuencia (por ejemplo, la -s de tercera persona, artículos o in / on / at)."],
  app="Sol: por qué se retrasa algo y qué pasa si no se resuelve. Virginia: el impacto de una decisión. Belén: qué pasa si el proveedor no completa un paso.",
  retry="La misma explicación. Presión: «What if that doesn't work?».",
  next="Registrar su patrón de alta frecuencia en cada práctica de la semana.",
  ya="Que proponga una solución con condiciones."),

 dict(n=16, sem=16, etapa="Polish Your English", nombre="Integración de precisión",
  core="Sostener la precisión en una situación real completa, bajo presión.",
  plat="Repaso de M3 (lo que corresponda a sus patrones).",
  res="Mantiene estables sus patrones prioritarios mientras habla en tiempo real, no solo cuando piensa en la gramática.",
  ev="Comparación de la situación de la sesión 10 con la de hoy · lista actualizada: qué ya está resuelto y qué queda.",
  hito="Patrón prioritario estable. Si todavía no lo está, sigue como track transversal en las sesiones 17–19.",
  foco="sostener tu precisión cuando hablas en tiempo real", puede="hablar con precisión sin frenar",
  micro=["No hay contenido nuevo: repaso de sus 2–3 patrones en tres minutos.",
         "Autocorrección rápida: detectar, corregir y seguir."],
  app="Repite la situación del hito de la sesión 10, o una equivalente, completa.",
  retry="Una segunda vuelta más rápida. Presión: interrupciones y preguntas.",
  next="Sus patrones pasan a ser un track transversal: los registra en el Weekly Report.",
  ya="—"),

 dict(n=17, sem=17, etapa="Understand & Be Understood", nombre="Pronunciación para ser entendida",
  core="Hacerse entender: acento de palabra, énfasis en lo importante, pausas y los sonidos que cambian el significado.",
  plat="M4.2 Pronunciation for Intelligibility · M4.1 How English Really Sounds.",
  res="Se hace entender sin esfuerzo del oyente en las palabras clave de su trabajo y marca el mensaje principal con énfasis y pausas.",
  ev="Diez palabras clave de su trabajo con el acento correcto · update grabado con pausas por bloques de sentido y énfasis en la información nueva.",
  foco="que te entiendan sin esfuerzo", puede="marcar tu mensaje con énfasis y pausas",
  micro=["El objetivo es ser entendida, no sonar nativa.",
         "Word stress en sus palabras clave. Sentence stress: énfasis en lo nuevo o importante.",
         "Pausas por bloques de sentido.",
         "Los 1–2 sonidos que, en su caso, cambian el significado (-ed, -s, vocales largas y cortas…)."],
  app="Sus palabras clave. Belén: vendor, compliance, onboarding. Sol: deliverable, timeline, schedule. Virginia: revenue, forecast, quarterly.",
  retry="El mismo update, con pausas y énfasis marcados. Presión: a velocidad normal de reunión.",
  next="Shadowing de un minuto de su propio update corregido. Grabarse.",
  ya="Trabajar entonación en preguntas y en recomendaciones."),

 dict(n=18, sem=18, etapa="Understand & Be Understood", nombre="Entrenar el oído",
  core="Entender inglés hablado real: formas débiles, uniones, reducciones y acentos distintos.",
  plat="M4.3 Train Your Ear · M4.1 How English Really Sounds.",
  res="Entiende la idea principal y los detalles clave de un audio real de su área, y reconoce formas reducidas.",
  ev="Resumen oral de un audio real (idea principal + dos detalles) · dictado corto de frases con connected speech.",
  foco="entender el inglés como se habla de verdad", puede="seguir una conversación real sin perder lo importante",
  micro=["Por qué no entiende aunque conozca las palabras: formas débiles (to, for, and, of), uniones (turn it off), reducciones (gonna, wanna, kinda).",
         "Primero la idea, después el detalle.",
         "En una reunión: confirmar lo que entendió en lugar de pedir que repitan todo."],
  app="Un audio de su contexto: la grabación de una reunión real (si puede compartirla) o un video de su industria, con los acentos con los que trabaja.",
  retry="Segunda escucha para completar lo que faltó; después responde como si estuviera en la reunión.",
  next="Diez minutos diarios de escucha de su área + un resumen de 30 segundos.",
  ya="Audios más rápidos o con acentos que le cuestan."),

 dict(n=19, sem=19, etapa="Understand & Be Understood", nombre="Escuchar y responder en tiempo real",
  core="Escuchar, procesar y responder en tiempo real; shadowing para ritmo y fluidez.",
  plat="M4.4 Shadowing.",
  res="Escucha una intervención real, confirma lo que entendió y responde con estructura, sin pedir que le repitan todo.",
  ev="Role play de reunión donde escucha, parafrasea («So what you're saying is…») y responde con estructura · audio de shadowing antes y después.",
  foco="escuchar y responder en el momento", puede="participar en una reunión real con más seguridad",
  micro=["Shadowing: escuchar y repetir en simultáneo imitando ritmo y entonación, en tres rondas.",
         "Escuchar → confirmar → responder.",
         "Chunks: «If I understand correctly…», «Just to make sure…»."],
  app="Reunión simulada de su contexto. Virginia recibe feedback del CEO y responde. Sol recibe un cambio de prioridad. Belén: el proveedor pone una objeción.",
  retry="La misma reunión. Presión: el interlocutor habla rápido o con otro acento.",
  next="Preparar la sesión 20: elegir la situación final y volver a grabar la consigna del Starting Point.",
  ya="Dos interlocutores que no están de acuerdo."),

 dict(n=20, sem=20, etapa="Offboarding", nombre="Offboarding · review final e integración",
  core="Integrar todo y medir el cambio con evidencia.",
  plat="Starting Point + carpeta 05 · MI PROGRESO + recurso «Tu recorrido AT WORK».",
  res="Compara su punto de partida con el actual, reconoce lo que ya hace sola y define su siguiente nivel y su plan de práctica autónoma.",
  ev="Starting Point vs producción final (misma consigna, 2–3 min) · situación real final completa · plan de práctica autónoma.",
  foco="mirar todo tu recorrido y definir cómo sigues", puede="seguir practicando sola, con método",
  micro=["No hay contenido nuevo: cómo seguir sola (ciclo de práctica, bots, su propio Weekly Report)."],
  app="Una situación real final completa (update + preguntas) usando todo: estructura, precisión e inteligibilidad. Después escuchan juntas el audio del Starting Point y el de hoy.",
  retry="Una sola vuelta final: su mejor versión.",
  next="Plan de práctica autónoma y su siguiente nivel: qué seguir entrenando después de AT WORK.",
  ya="—"),
]

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2.2); sec.top_margin = sec.bottom_margin = Cm(2)

st = doc.styles["Normal"]; st.font.name = BODY; st.font.size = Pt(10.5); st.font.color.rgb = RGBColor.from_string(DARK)
st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY)
st.paragraph_format.space_after = Pt(4); st.paragraph_format.line_spacing = 1.2

def run(p, text, bold=False, italic=False, size=None, color=None, font=None):
    r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    if font: r.font.name = font; r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    return r

def para(text="", **kw):
    p = doc.add_paragraph()
    if text: run(p, text, **kw)
    return p

def shade(cell, color):
    tcPr = cell._element.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), color); tcPr.append(sh)

def borders(table, color="EDE5CC"):
    tblPr = table._element.tblPr; b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), color); b.append(el)
    anchor = tblPr.find(qn("w:tblLayout"))
    if anchor is None: anchor = tblPr.find(qn("w:tblLook"))
    anchor.addprevious(b) if anchor is not None else tblPr.append(b)

def widths(t, ws):
    t.autofit = False
    grid = t._element.find(qn("w:tblGrid"))
    for gc, w in zip(grid.findall(qn("w:gridCol")), ws): gc.set(qn("w:w"), str(int(Cm(w).twips)))
    for row in t.rows:
        for c, w in zip(row.cells, ws): c.width = Cm(w)

def label(text):
    p = para(); p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
    run(p, text.upper(), bold=True, size=8.5, color=TERRA)

def box(text, fill=TINT, lab=None):
    t = doc.add_table(rows=1, cols=1); borders(t, fill); c = t.cell(0, 0); shade(c, fill)
    p = c.paragraphs[0]
    if lab: run(p, lab + "  ", bold=True, size=9, color=BROWN)
    run(p, text, size=10)
    para().paragraph_format.space_after = Pt(0)

# Portada
p = para(); run(p, "BREAKING BARRIERS AT WORK", bold=True, size=9, color=TERRA)
p = para(); run(p, "Master Delivery Map", bold=True, size=28, color=BROWN, font=TITLE)
p = para(); run(p, "20 sesiones 1:1 · la columna vertebral del programa", italic=True, size=14, color=MUTED, font=TITLE)
p = para(); run(p, "Borrador para revisar antes de subir a Notion · by Flor Pérsico", size=9, color=MUTED)

label("Cómo leer este documento")
para("Cada alumna tiene una sesión 1:1 por semana: 20 sesiones en total. La sesión 1 es su onboarding, las sesiones 2–19 recorren el método completo y la sesión 20 es el offboarding.")
para("La plataforma es soporte. La sesión 1:1 garantiza que cada alumna pase por todas las etapas y todas las habilidades esenciales, aunque no haya visto un video.")
t = doc.add_table(rows=1, cols=2); borders(t)
for i, (h, b) in enumerate([("LO FIJO", "Etapa · Core Skill · resultado esperado · evidencia observable · microenseñanza"),
                             ("LO QUE VARÍA", "Su situación real: lo que trae en el Weekly Report o lo que tiene esa semana en el trabajo. Si el Weekly Report pide una situación concreta, esa situación manda; la habilidad se mantiene.")]):
    c = t.cell(0, i); shade(c, TINT if i == 0 else CARD); cp = c.paragraphs[0]
    run(cp, h + "\n", bold=True, size=9, color=BROWN); run(cp, b, size=10)
para()

label("Estructura fija de cada sesión")
steps = [("Apertura", "Etapa + foco + resultado: «Estamos en…, hoy el foco es…, al terminar vas a poder…»."),
         ("Situación real y primer intento", "Produce antes de recibir ayuda. Sin frameworks ni frases antes del primer intento."),
         ("Microenseñanza (5–10 min)", "Solo si no vio la clase de la plataforma o si el primer intento muestra que no la incorporó. Si ya la vio: una pregunta de chequeo y directo a la aplicación."),
         ("Aplicación a su situación real", "La herramienta aplicada a su trabajo, no a un ejemplo inventado."),
         ("Retry + presión", "La misma tarea otra vez para comparar intento 1 vs intento 2; después, algo de imprevisibilidad."),
         ("Cierre", "Evidencia + gap + siguiente acción: qué logró hoy, qué falta y qué hace esta semana."),
         ("Next step", "Continúa exactamente la habilidad de la sesión. Nada de tarea desconectada.")]
t = doc.add_table(rows=len(steps), cols=2); borders(t)
for i, (a, b) in enumerate(steps):
    c0, c1 = t.cell(i, 0), t.cell(i, 1); shade(c0, TINT)
    run(c0.paragraphs[0], f"{i+1} · {a}", bold=True, size=10, color=BROWN); run(c1.paragraphs[0], b, size=10)
widths(t, [5, 11.6])
para()
box("Si una alumna todavía no tiene autonomía en una habilidad, la sesión siguiente introduce igual la nueva, y la pendiente sigue como track transversal. Así nadie queda sin pasar por una habilidad esencial y el recorrido no se corre.", lab="REGLA")

label("Mapa general")
etapas = [("Semana 1", "Onboarding", "Sesión 1"), ("Semanas 2–4", "Activate Your English", "Sesiones 2–4 · hito en la 4"),
          ("Semanas 5–10", "Communicate in Real Time", "Sesiones 5–10 · hito en la 10"), ("Semanas 11–16", "Polish Your English", "Sesiones 11–16 · hito en la 16"),
          ("Semanas 17–19", "Understand & Be Understood", "Sesiones 17–19"), ("Semana 20", "Offboarding · review final", "Sesión 20")]
t = doc.add_table(rows=len(etapas), cols=3); borders(t)
for i, row in enumerate(etapas):
    for j, v in enumerate(row):
        c = t.cell(i, j); run(c.paragraphs[0], v, bold=(j == 1), size=10, color=BROWN if j == 1 else DARK)
        if j == 0: shade(c, TINT)
widths(t, [3.6, 6.5, 6.5])
para("Tracks transversales durante todo el recorrido: práctica con situaciones reales, pronunciación y listening, y precisión.", italic=True, size=9.5, color=MUTED)

# Sesiones
for s in S:
    doc.add_page_break()
    p = para(); run(p, f"SESIÓN {s['n']} · SEMANA {s['sem']} · {s['etapa'].upper()}", bold=True, size=9, color=TERRA)
    p = para(); run(p, s["nombre"], bold=True, size=20, color=BROWN, font=TITLE)
    rows = [("Core Skill", s["core"]), ("Apoyo en plataforma", s["plat"]), ("Resultado esperado", s["res"]), ("Evidencia observable", s["ev"])]
    if s.get("hito"): rows.append(("Hito", s["hito"]))
    t = doc.add_table(rows=len(rows), cols=2); borders(t)
    for i, (a, b) in enumerate(rows):
        c0, c1 = t.cell(i, 0), t.cell(i, 1); shade(c0, TINT if a != "Hito" else "F3E3D8")
        run(c0.paragraphs[0], a, bold=True, size=9.5, color=BROWN); run(c1.paragraphs[0], b, size=10)
    widths(t, [4.2, 12.4])
    label("Apertura")
    para(f"«Estamos en {s['etapa'] if s['etapa'] not in ('Onboarding','Offboarding') else 'tu ' + s['etapa'].lower()}. Hoy el foco es {s['foco']}. Al terminar, vas a poder {s['puede']}.»", italic=True, color=MUTED)
    label("Microenseñanza · 5–10 min (si no vio la plataforma)")
    for m in s["micro"]:
        doc.add_paragraph(m, style="List Bullet")
    label("Aplicación a su situación real")
    para(s["app"])
    label("Retry + presión")
    para(s["retry"])
    label("Next step")
    para(s["next"])
    if s["ya"] != "—":
        label("Si ya lo domina")
        para(s["ya"], italic=True, color=MUTED)

out = "recursos/Master_Delivery_Map_20_sesiones.docx"
doc.save(out); print(out)
