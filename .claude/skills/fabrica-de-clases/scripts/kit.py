"""Arma deck (.pptx con el prompter en las notas) y prompter (.docx y .txt) desde un spec JSON.

Uso: python3 kit.py spec.json carpeta_salida

Marca copiada del deck M2.07 (Presentaciones v6): portada marrón, slides crema,
barra superior cobre, rótulo Arial y título Georgia.
"""
import json
import sys
from pathlib import Path

import docx
from docx.shared import Pt, RGBColor as DocRGB
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt as PPt

MARRON = RGBColor(0x3F, 0x1E, 0x0D)
CREMA = RGBColor(0xF9, 0xF5, 0xEE)
COBRE = RGBColor(0xC6, 0x8C, 0x6C)
TOSTADO = RGBColor(0x89, 0x53, 0x25)
GRIS = RGBColor(0x7B, 0x6A, 0x58)


def caja(slide, x, y, w, h, lineas):
    """lineas: lista de (texto, fuente, tamaño, negrita, color)."""
    tf = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h)).text_frame
    tf.word_wrap = True
    for i, (texto, fuente, tam, negrita, color) in enumerate(lineas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = texto
        r.font.name, r.font.size, r.font.bold = fuente, PPt(tam), negrita
        r.font.color.rgb = color
        p.space_after = PPt(10 if fuente == "Georgia" else 2)
    return tf


def fondo(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def deck(spec, salida):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12191695), Emu(6858000)
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    fondo(s, MARRON)
    caja(s, 3840480, 1463040, 7772400, 457200, [(spec["etiqueta"], "Arial", 14, True, COBRE)])
    caja(s, 3840480, 2011680, 7772400, 1463040, [(spec["titulo"], "Georgia", 60, True, CREMA)])
    caja(s, 3840480, 3566160, 7772400, 1280160,
         [(l, "Georgia", 28, False, COBRE) for l in spec["subtitulo"]])
    s.notes_slide.notes_text_frame.text = spec["portada_notas"]

    for sl in spec["slides"]:
        s = prs.slides.add_slide(blank)
        fondo(s, CREMA)
        barra = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Emu(109728))
        barra.fill.solid()
        barra.fill.fore_color.rgb = COBRE
        barra.line.fill.background()
        caja(s, 548640, 320040, 10972800, 365760, [(sl["rotulo"], "Arial", 13, True, TOSTADO)])
        caja(s, 548640, 685800, 10972800, 1188720,
             [(l, "Georgia", 34, True, MARRON) for l in sl["titulo"].split("\n")])
        cuerpo = []
        for item in sl["cuerpo"]:
            if item.startswith("# "):
                cuerpo.append((item[2:], "Arial", 12, True, TOSTADO))
            else:
                cuerpo.append((item, "Georgia", 18, False, MARRON))
        caja(s, 3840480, 1920240, 7863840, 4663440, cuerpo)
        if sl.get("lateral"):
            caja(s, 548640, 1920240, 2926080, 2194560, [(sl["lateral"], "Georgia", 14, False, GRIS)])
        s.notes_slide.notes_text_frame.text = sl["notas"]

    prs.save(salida)


def bloques(spec):
    yield 1, spec["portada_notas"]
    for i, sl in enumerate(spec["slides"], start=2):
        yield i, sl["notas"]


def prompter(spec, carpeta, base):
    nombre = f'{spec["codigo"]} · {spec["titulo"]}'
    d = docx.Document()
    h = d.add_paragraph()
    r = h.add_run(nombre)
    r.bold, r.font.size = True, Pt(18)
    sub = d.add_paragraph().add_run("PROMPTER LISTO · español neutro · igual a las notas de la presentación")
    sub.font.size, sub.font.color.rgb = Pt(10), DocRGB(0xC6, 0x8C, 0x6C)
    txt = [nombre, ""]
    palabras = 0
    for n, notas in bloques(spec):
        d.add_paragraph().add_run(f"SLIDE {n}").bold = True
        txt.append(f"SLIDE {n}")
        for parrafo in notas.split("\n"):
            if parrafo.strip():
                d.add_paragraph(parrafo).runs[0].font.size = Pt(14)
                txt.append(parrafo)
                palabras += len(parrafo.split())
        txt.append("")
    d.save(carpeta / f"{base} · PROMPTER LISTO.docx")
    (carpeta / f"{base} · PROMPTER.txt").write_text("\n".join(txt), encoding="utf-8")
    return palabras


def main():
    spec = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    carpeta = Path(sys.argv[2])
    carpeta.mkdir(parents=True, exist_ok=True)
    base = f'{spec["codigo"]} — {spec["titulo"]}'
    deck(spec, carpeta / f"{base} · DECK.pptx")
    palabras = prompter(spec, carpeta, base)
    print(f"{len(spec['slides']) + 1} slides · {palabras} palabras · ~{palabras / 130:.0f} min")


if __name__ == "__main__":
    main()
